"""
Safe text extraction for uploaded CVs / JDs (PDF, DOCX, TXT).

Every upload is untrusted. Before any parser touches the bytes we check:
  * size limit                      (MAX_FILE_BYTES)
  * magic bytes match the extension (a renamed .exe is not a PDF)
  * page limit for PDFs             (MAX_PDF_PAGES — parser DoS guard)
  * zip-bomb guard for DOCX         (MAX_DOCX_UNCOMPRESSED)

Scanned (image-only) PDF pages are OCR'd with RapidOCR (ONNX, fully
local) when it is installed. Previously such CVs were silently
rejected as "unreadable".

`extract_text()` keeps its old signature for callers that only need a
string; `extract_document()` returns the full ExtractionResult with
warnings and OCR metadata.
"""
import io
import logging
import re
import threading
import unicodedata
import zipfile
from dataclasses import dataclass, field

import pdfplumber

from config import (
    MAX_DOCX_UNCOMPRESSED, MAX_FILE_BYTES, MAX_OCR_PAGES, MAX_PDF_PAGES,
    MIN_CHARS_PER_PAGE_BEFORE_OCR, OCR_ENABLED, OCR_RENDER_DPI,
)

logger = logging.getLogger(__name__)

SUPPORTED_TYPES = ("pdf", "docx", "txt")


class ExtractionError(Exception):
    """Raised with a user-safe message when a file cannot be read."""


@dataclass
class ExtractionResult:
    text: str
    file_type: str
    pages: int = 0
    ocr_pages: int = 0
    warnings: list = field(default_factory=list)


_ZERO_WIDTH_RE = re.compile(r'[​-‏⁠﻿­]')
_QUOTES = str.maketrans({"\u2018": "'", "\u2019": "'", "\u201b": "'", "\u2032": "'",
                        "\u201c": '"', "\u201d": '"', "\u2212": "-", "\u2010": "-", "\u2011": "-"})
_CTRL_RE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]')


def normalize_text(text: str) -> str:
    """
    NFKC folds ligatures ("ﬁ" → "fi", which broke "certiﬁcation"),
    full-width and math-styled letters ("𝐒𝐤𝐢𝐥𝐥𝐬" → "Skills") into plain
    text. Zero-width / control characters are removed, NBSP → space.
    """
    if not text:
        return ""
    text = unicodedata.normalize("NFKC", text)
    text = _ZERO_WIDTH_RE.sub("", text)
    text = text.translate(_QUOTES)
    text = _CTRL_RE.sub(" ", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r'[ \t ]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


def _read_bytes(file) -> bytes:
    if isinstance(file, (bytes, bytearray)):
        return bytes(file)
    if hasattr(file, "getvalue"):
        return file.getvalue()
    if hasattr(file, "seek"):
        try:
            file.seek(0)
        except Exception:
            pass
    return file.read()


def detect_type(filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return ext


def _validate(data: bytes, file_type: str) -> None:
    if file_type not in SUPPORTED_TYPES:
        raise ExtractionError(f"Unsupported file type '.{file_type}'. Use PDF, DOCX or TXT.")
    if not data:
        raise ExtractionError("File is empty (0 bytes).")
    if len(data) > MAX_FILE_BYTES:
        raise ExtractionError(
            f"File is too large ({len(data) / 1_048_576:.1f} MB; limit "
            f"{MAX_FILE_BYTES / 1_048_576:.0f} MB).")
    if file_type == "pdf" and b"%PDF-" not in data[:1024]:
        raise ExtractionError("Not a valid PDF (file signature missing). Was the file renamed?")
    if file_type == "docx" and not data.startswith(b"PK"):
        raise ExtractionError("Not a valid DOCX (legacy .doc or renamed file?). Save it as .docx.")
    if file_type == "txt" and b"\x00" in data[:4096] and data[:2] not in (b"\xff\xfe", b"\xfe\xff"):
        raise ExtractionError("File looks binary, not plain text.")


def extract_document(file, file_type: str) -> ExtractionResult:
    """Validate and extract. Raises ExtractionError with a user-safe message."""
    file_type = (file_type or "").lower().lstrip(".")
    data = _read_bytes(file)
    _validate(data, file_type)

    try:
        if file_type == "pdf":
            result = _extract_pdf(data)
        elif file_type == "docx":
            result = _extract_docx(data)
        else:
            result = ExtractionResult(_decode_txt(data), "txt")
    except ExtractionError:
        raise
    except Exception as exc:
        logger.warning("extraction failed: %s", type(exc).__name__)
        raise ExtractionError(
            f"File is corrupted or could not be parsed ({type(exc).__name__}).") from exc

    result.text = normalize_text(result.text)
    return result


def extract_text(file, file_type: str) -> str:
    """Backward-compatible helper: text only."""
    return extract_document(file, file_type).text


def _decode_txt(data: bytes) -> str:
    """UTF-8/UTF-16 by BOM, then UTF-8, then charset detection, then cp1252."""
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    try:
        from charset_normalizer import from_bytes
        best = from_bytes(data).best()
        if best is not None:
            return str(best)
    except Exception:
        pass
    return data.decode("cp1252", errors="replace")


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_body_lines(body) -> list:
    """Text of every paragraph (w:p) under body, grouped per paragraph, in
    document order. Text-box content appears twice in the XML
    (mc:Choice and mc:Fallback) — the fallback copy is skipped."""
    lines = {}
    for node in body.iter(_W + "t", _W + "tab", _W + "br"):
        para = None
        anc = node.getparent()
        skip = False
        while anc is not None:
            tag = anc.tag if isinstance(anc.tag, str) else ""
            if tag.endswith("}Fallback"):
                skip = True
                break
            if para is None and tag == _W + "p":
                para = anc
            anc = anc.getparent()
        if skip or para is None:
            continue
        bucket = lines.setdefault(para, [])
        if node.tag == _W + "t":
            bucket.append(node.text or "")
        elif node.tag == _W + "tab":
            bucket.append(" ")
        else:
            bucket.append("\n")
    return ["".join(parts) for parts in lines.values()]


def _extract_docx(data: bytes) -> ExtractionResult:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            total = sum(i.file_size for i in zf.infolist())
            if total > MAX_DOCX_UNCOMPRESSED:
                raise ExtractionError("DOCX expands to an unsafe size (possible zip bomb).")
            if "word/document.xml" not in zf.namelist():
                raise ExtractionError("Not a Word document (word/document.xml missing).")
    except zipfile.BadZipFile as exc:
        raise ExtractionError("DOCX is corrupted (bad zip structure).") from exc

    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = []

    for section in doc.sections:
        for p in section.header.paragraphs:
            if p.text.strip():
                parts.append(p.text)

    parts.extend(_docx_body_lines(doc.element.body))
    return ExtractionResult("\n".join(parts), "docx", pages=0)


def _group_lines(w_list):
    if not w_list:
        return ""
    w_list = sorted(w_list, key=lambda w: (w['top'], w['x0']))
    lines = []
    curr_line = [w_list[0]]
    for w in w_list[1:]:
        if abs(w['top'] - curr_line[-1]['top']) <= 4:
            curr_line.append(w)
        else:
            curr_line.sort(key=lambda x: x['x0'])
            lines.append(" ".join(x['text'] for x in curr_line))
            curr_line = [w]
    curr_line.sort(key=lambda x: x['x0'])
    lines.append(" ".join(x['text'] for x in curr_line))
    return "\n".join(lines)


def _page_text_columns(page) -> str:
    """Reading-order text for one page, splitting two-column layouts."""
    words = page.extract_words(x_tolerance=2, y_tolerance=3)
    if not words:
        return page.extract_text() or ""

    pw = page.width
    mid_start = int(pw * 0.20)
    mid_end = int(pw * 0.80)
    word_boxes = [(w['x0'], w['x1']) for w in words]

    best_split = None
    best_gutter_width = 0
    for x in range(mid_start, mid_end, 5):
        if any(not (b[1] < x - 5 or b[0] > x + 5) for b in word_boxes):
            continue
        left_edge = max((b[1] for b in word_boxes if b[1] <= x), default=0)
        right_edge = min((b[0] for b in word_boxes if b[0] >= x), default=pw)
        gutter = right_edge - left_edge
        if gutter >= 10 and gutter > best_gutter_width:
            left_cnt = sum(1 for b in word_boxes if b[1] <= x)
            right_cnt = sum(1 for b in word_boxes if b[0] >= x)
            if left_cnt >= len(words) * 0.15 and right_cnt >= len(words) * 0.15:
                best_gutter_width = gutter
                best_split = (left_edge + right_edge) / 2

    if best_split is None:
        return _group_lines(words)

    left_words = [w for w in words if w['x1'] <= best_split + 5]
    right_words = [w for w in words if w['x0'] >= best_split - 5 and w['x1'] > best_split + 5]
    middle_words = [w for w in words if w['x0'] < best_split - 5 and w['x1'] > best_split + 5]
    first_col_top = min(
        min((w['top'] for w in left_words), default=9999),
        min((w['top'] for w in right_words), default=9999),
    )
    top_span = [w for w in middle_words if w['top'] < first_col_top]
    other_middle = [w for w in middle_words if w['top'] >= first_col_top]

    page_lines = []
    for chunk in (top_span, left_words, right_words, other_middle):
        if chunk:
            page_lines.append(_group_lines(chunk))
    return "\n".join(page_lines)


_ocr_engine = None
_ocr_lock = threading.Lock()


def _get_ocr():
    """Lazy singleton — RapidOCR loads ~3 ONNX models (~1 s). Call with _ocr_lock held."""
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr_engine = RapidOCR()
    return _ocr_engine


def ocr_available() -> bool:
    if not OCR_ENABLED:
        return False
    try:
        import pypdfium2  # noqa: F401
        import rapidocr_onnxruntime  # noqa: F401
        return True
    except ImportError:
        return False


def _ocr_page(pdf_bytes: bytes, page_index: int) -> str:
    import numpy as np
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(pdf_bytes)
    try:
        page = doc[page_index]
        bitmap = page.render(scale=OCR_RENDER_DPI / 72)
        img = bitmap.to_pil().convert("RGB")
        with _ocr_lock:
            result, _ = _get_ocr()(np.array(img))
    finally:
        doc.close()
    if not result:
        return ""
    rows = sorted(((b[0][1], b[0][0], t) for b, t, conf in result if float(conf) >= 0.5))
    lines, current, last_y = [], [], None
    for y, x, t in rows:
        if last_y is not None and abs(y - last_y) > 12:
            lines.append(" ".join(s for _, s in sorted(current)))
            current = []
        current.append((x, t))
        last_y = y
    if current:
        lines.append(" ".join(s for _, s in sorted(current)))
    return "\n".join(lines)


def _extract_pdf(data: bytes) -> ExtractionResult:
    warnings = []
    texts = []
    ocr_pages = 0
    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as exc:
        msg = str(exc).lower()
        if "password" in msg or "encrypt" in msg:
            raise ExtractionError("PDF is password-protected.") from exc
        raise ExtractionError(f"PDF is corrupted ({type(exc).__name__}).") from exc

    with pdf:
        n_pages = len(pdf.pages)
        if n_pages == 0:
            raise ExtractionError("PDF has no pages.")
        if n_pages > MAX_PDF_PAGES:
            warnings.append(f"Only the first {MAX_PDF_PAGES} of {n_pages} pages were read.")
        can_ocr = ocr_available()
        for i, page in enumerate(pdf.pages[:MAX_PDF_PAGES]):
            try:
                page_text = _page_text_columns(page)
            except Exception as exc:
                logger.warning("page %d text extraction failed: %s", i, type(exc).__name__)
                page_text = ""
            has_image = bool(getattr(page, "images", None))
            if len(page_text.strip()) < MIN_CHARS_PER_PAGE_BEFORE_OCR and can_ocr and has_image:
                if ocr_pages < MAX_OCR_PAGES:
                    try:
                        ocr_text = _ocr_page(data, i)
                        if len(ocr_text.strip()) > len(page_text.strip()):
                            page_text = ocr_text
                            ocr_pages += 1
                    except Exception as exc:
                        logger.warning("OCR failed on page %d: %s", i, type(exc).__name__)
                        warnings.append(f"OCR failed on page {i + 1}.")
                else:
                    warnings.append(f"Page {i + 1} looks scanned but the OCR page limit was reached.")
            texts.append(page_text)

    if can_ocr and not ocr_pages and sum(len(t.strip()) for t in texts) < MIN_CHARS_PER_PAGE_BEFORE_OCR:
        for i in range(min(len(texts), MAX_OCR_PAGES)):
            try:
                ocr_text = _ocr_page(data, i)
            except Exception as exc:
                logger.warning("OCR failed on page %d: %s", i, type(exc).__name__)
                continue
            if ocr_text.strip():
                texts[i] = ocr_text
                ocr_pages += 1

    if ocr_pages:
        warnings.append(f"{ocr_pages} scanned page(s) read with OCR — check extracted details.")
    elif not any(t.strip() for t in texts) and not can_ocr:
        warnings.append("No text layer found and OCR is not installed (scanned PDF?).")
    return ExtractionResult("\n\n".join(texts), "pdf", pages=n_pages,
                            ocr_pages=ocr_pages, warnings=warnings)
