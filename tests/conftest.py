import glob
import io
import os
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

SAMPLES = os.path.join(ROOT, "CV Project Test Cases")


def sample_cvs(folder=None):
    pattern = os.path.join(SAMPLES, folder or "*", "*.pdf")
    return sorted(glob.glob(pattern))


def read(path, mode="rb"):
    with open(path, mode, **({} if "b" in mode else {"encoding": "utf-8"})) as f:
        return f.read()


@pytest.fixture(scope="session")
def jd_texts():
    return {
        "ai": read(os.path.join(SAMPLES, "AI JD.txt"), "r"),
        "fullstack": read(os.path.join(SAMPLES, "Full Stack JD.txt"), "r"),
        "mobile": read(os.path.join(SAMPLES, "AI Mobile App.txt"), "r"),
    }


@pytest.fixture(scope="session")
def all_cv_files():
    found = sample_cvs()
    if not found:
        pytest.skip("Original private resume fixture PDFs deliberately excluded from distributable source")
    return [(os.path.relpath(p, SAMPLES), read(p)) for p in found]


def make_docx(paragraphs=(), table_rows=None, header=None) -> bytes:
    from docx import Document
    d = Document()
    if header:
        d.sections[0].header.paragraphs[0].text = header
    for p in paragraphs:
        d.add_paragraph(p)
    if table_rows:
        t = d.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for i, row in enumerate(table_rows):
            for j, val in enumerate(row):
                t.cell(i, j).text = val
        d.add_paragraph("AFTER TABLE")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_image_pdf(pages=1, text_pdf_page_source=None, scale=2.0) -> bytes:
    """Image-only PDF. With text_pdf_page_source, page 1 is a raster of that PDF."""
    from PIL import Image
    imgs = []
    if text_pdf_page_source:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(text_pdf_page_source)
        imgs.append(doc[0].render(scale=scale).to_pil().convert("RGB"))
        doc.close()
    while len(imgs) < pages:
        imgs.append(Image.new("RGB", (200, 260), "white"))
    buf = io.BytesIO()
    imgs[0].save(buf, format="PDF", save_all=True, append_images=imgs[1:])
    return buf.getvalue()
