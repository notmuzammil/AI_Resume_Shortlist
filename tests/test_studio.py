"""Local direct-launch recruiter product integration and safety tests (no Streamlit dependency)."""
import csv
import io
import json
import zipfile

import pytest

from studio_core import (
    build_report, clean_name, csv_escape, ensure_valid_cv, export_csv, export_json,
    experience_meter, normalize_category, report_statistics,
    unpack_files,
)
from studio_store import (
    all_counts, candidate_file, connect, create_job, delete_candidate, delete_job,
    delete_run, get_job, get_run, initialize, insert_candidate, list_candidates,
    list_jobs, list_runs, save_run, update_job,
)
from modules.pipeline import evaluate_batch
from modules.extractor import ExtractionError

JD='''Job Title: Senior Backend Engineer\nRequired skills: Node.js, PostgreSQL, Docker.\nMinimum 3 years of experience.\nA completed bachelor's degree in Computer Science is required.\n'''
GOOD='''Jane Developer\nProfile\nSenior software engineer with Node.js, PostgreSQL and Docker in production.\nEducation\nBS Computer Science (2017)\nExperience\nSenior Backend Engineer | January 2018 - December 2024\nBuilt Node.js and PostgreSQL microservices with Docker for enterprise API integration.\n'''.encode()
BAD='''J. Designer\nGraphic designer with a focus on Figma, Illustrator and branding.\nEducation\nBA Graphic Design (2018)\nExperience\nGraphic Designer | Jan 2019 - Dec 2023\nDesigned marketing banners and logos for agencies using Adobe Photoshop.\n'''.encode()

@pytest.fixture
def workspace(tmp_path,monkeypatch):
    monkeypatch.setenv('RESUME_STUDIO_DATA_DIR',str(tmp_path/'local'))
    initialize()
    yield


def test_direct_local_create_job_update_delete(workspace):
    jid=create_job('Backend Role',JD)
    assert get_job(jid)['description']==JD.strip()
    assert len(list_jobs())==1
    update_job(jid,'New Backend Title',JD+'\nPreferred: AWS.')
    assert get_job(jid)['title']=='New Backend Title'
    delete_job(jid)
    assert not list_jobs()


def test_jobs_validate_jd_and_name(workspace):
    with pytest.raises(ValueError):create_job('',JD)
    with pytest.raises(ValueError):create_job('Short','too short')
    with pytest.raises(ValueError):update_job('missing','Name',JD)


def test_candidate_dedupe_and_roundtrip(workspace):
    a,created=insert_candidate('Jane.txt',GOOD)
    assert created
    again,created=insert_candidate('Jane-copy.txt',GOOD)
    assert not created and again['id']==a['id']
    assert candidate_file(a['id'])['content']==GOOD
    assert all_counts()['candidates']==1
    assert list_candidates()[0]['filename']=='Jane.txt'


def test_e2e_run_rank_score_and_json_csv(workspace):
    jid=create_job('Backend',JD)
    a,_=insert_candidate('Jane.txt',GOOD)
    b,_=insert_candidate('Graphic.txt',BAD)
    uploaded=[candidate_file(x['id']) for x in (a,b)]
    result=evaluate_batch(JD,[(x['filename'],x['content']) for x in uploaded])
    report=build_report(result,jid,'Backend',{x['sha256']:x['id'] for x in uploaded})
    assert report['statistics']['processed']==2
    assert report['statistics']['scored']>=1
    scored=[x for x in report['candidates'] if x['status']=='scored']
    assert scored[0]['filename']=='Jane.txt'
    assert report['statistics']['highly_suitable'] == sum(x['category']=='Highly Suitable' for x in scored)
    assert scored[0]['components']['skills']['percent'] is not None
    assert all('match_percent' in r and 'evidence' in r for r in scored[0]['requirements'])
    assert all(0<=r['match_percent']<=100 for r in scored[0]['requirements'])
    assert scored[0]['relevant_experience']['percent'] <= 100
    assert scored[0]['total_experience']['percent'] <= 100
    assert scored[0]['relevant_years']<=scored[0]['total_years']
    rid=save_run(jid,report)
    assert get_run(rid)['job_title']=='Backend'
    assert len(list_runs(jid))==1
    raw=export_json(report)
    assert json.loads(raw)['candidates'][0]['candidate_id']==a['id']
    rows=list(csv.DictReader(io.StringIO(export_csv(report).decode('utf-8-sig'))))
    assert len(rows)==2 and rows[0]['match_percent']!=''
    delete_run(rid)
    assert get_run(rid) is None


def test_delete_candidate_scrubs_saved_evidence(workspace):
    jid=create_job('Backend',JD)
    a,_=insert_candidate('Jane.txt',GOOD)
    res=evaluate_batch(JD,[('Jane.txt',GOOD)])
    rid=save_run(jid,build_report(res,jid,'Backend',{res.candidates[0].sha256:a['id']}))
    assert get_run(rid)['candidates'][0]['candidate_id']==a['id']
    assert delete_candidate(a['id'])
    assert get_run(rid)['candidates']==[]
    assert b'Jane Developer' not in json.dumps(get_run(rid)).encode()
    assert all_counts()['candidates']==0


def test_delete_all_candidates_keeps_jobs_and_purges_file(workspace):
    from studio_store import db_file, delete_all_candidates
    jid=create_job('Backend',JD)
    insert_candidate('Jane.txt',GOOD)
    insert_candidate('Graphic.txt',BAD)
    save_run(jid,{'job_title':'Backend','candidates':[{'candidate_id':'x','filename':'Jane.txt'}]})
    assert delete_all_candidates()==2
    assert all_counts()=={'jobs':1,'candidates':0,'runs':0}
    assert get_job(jid) is not None
    raw=db_file().read_bytes()
    assert b'Jane Developer' not in raw and b'Graphic designer' not in raw


def test_reset_workspace_deletes_everything(workspace):
    from studio_store import db_file, reset_workspace
    jid=create_job('Backend',JD)
    insert_candidate('Jane.txt',GOOD)
    save_run(jid,{'job_title':'Backend','candidates':[]})
    assert reset_workspace()=={'runs':1,'candidates':1,'jobs':1}
    assert all_counts()=={'jobs':0,'candidates':0,'runs':0}
    assert b'Jane Developer' not in db_file().read_bytes()
    create_job('Again',JD)
    assert all_counts()['jobs']==1


def test_delete_job_cascades_runs_not_candidates(workspace):
    jid=create_job('Backend',JD)
    a,_=insert_candidate('Jane.txt',GOOD)
    rid=save_run(jid,{'job_title':'Backend','candidates':[]})
    delete_job(jid)
    assert get_run(rid) is None
    assert candidate_file(a['id']) is not None


def test_untrusted_filename_safe_for_local_display():
    assert clean_name('../../private/Jane.txt')=='Jane.txt'
    assert clean_name('C:\\users\\john\\Bob.pdf')=='Bob.pdf'
    assert '\n' not in clean_name('a\nb.txt')


def test_zip_batch_unpack_filters_paths_and_limits():
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w') as archive:
        archive.writestr('candidates/Jane.txt',GOOD)
        archive.writestr('../evil.txt',GOOD)
        archive.writestr('payload.exe',b'MZ')
        archive.writestr('candidates/Graphic.txt',BAD)
    files,errors=unpack_files([('bundle.zip',out.getvalue())])
    assert len(files)==2
    assert files[0][0]=='Jane.txt'
    assert any('unsafe' in e for e in errors)
    assert all(x[0].endswith('.txt') for x in files)


def test_zip_invalid_and_wrong_ext():
    files,errors=unpack_files([('broken.zip',b'bad'),('malware.exe',b'bin')])
    assert not files and len(errors)==2


def test_file_validation_rejects_renamed_pdf():
    with pytest.raises(ExtractionError):ensure_valid_cv('evil.pdf',b'not really a PDF')


def test_file_validation_accepts_txt():
    assert 'Jane Developer' in ensure_valid_cv('jane.txt',GOOD).text


def test_score_and_category_are_numerical_and_consistent():
    assert normalize_category('✅ Highly Suitable')=='Highly Suitable'
    assert normalize_category('🟢 Suitable')=='Suitable'
    assert normalize_category('🟡 Moderate')=='Moderate'
    assert normalize_category('❌ Not Suitable')=='Not Suitable'
    assert experience_meter(12,3)['percent']==100
    assert experience_meter(1,4)['percent']==25
    assert experience_meter(1,0)['reference_type'].startswith('display')


def test_json_has_neither_passwords_nor_auth_metadata(workspace):
    jid=create_job('Backend',JD)
    a,_=insert_candidate('Jane.txt',GOOD)
    res=evaluate_batch(JD,[('Jane.txt',GOOD)])
    report=build_report(res,jid,'Backend',{res.candidates[0].sha256:a['id']})
    assert not {'access_token','password','tenant_id','jwt'} & set(report)


def test_csv_formula_injection_is_neutralized(workspace):
    assert csv_escape('=HYPERLINK("bad")').startswith("'=")
    assert csv_escape('+cmd').startswith("'+")
    assert csv_escape('normal')=='normal'
    res=evaluate_batch(JD,[('=CALL.txt',GOOD)])
    report=build_report(res,'x','job',{})
    blob=export_csv(report).decode('utf-8-sig')
    assert "'=CALL.txt" in blob


def test_missing_requirement_statistics_are_derived_from_scored():
    items=[{'status':'scored','category':'Suitable','score':55.,'requirements':[
        {'skill':'Docker','matched':False},{'skill':'PostgreSQL','matched':True}]},
        {'status':'rejected','category':'Rejected','score':None,'requirements':[
        {'skill':'Docker','matched':False}]}]
    s=report_statistics(items)
    assert s['top_missing_skills']==[{'skill':'Docker','count':1}]
    assert s['rejected']==1 and s['average_score']==55


def test_private_local_db_does_not_use_auth(workspace):
    assert all_counts()=={'jobs':0,'candidates':0,'runs':0}
    with connect() as db:
        tables={r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert {'jobs','candidates','runs'} <= tables
    assert not {'users','sessions','tokens','tenants'} & tables


def test_pdf_page_viewer_real_pdf_is_rasterized():
    fitz=pytest.importorskip('fitz')
    from studio_core import render_pdf_page
    pdf=fitz.open()
    for i in (1,2):
        page=pdf.new_page()
        page.insert_text((70,90),f'Test candidate CV: 6 years of Node.js experience on page {i}')
    content=pdf.tobytes()
    data=ensure_valid_cv('candidate.pdf',content)
    assert data.pages==2
    png,count=render_pdf_page(content,2)
    assert count==2 and png.startswith(b'\x89PNG\r\n\x1a\n')
    with pytest.raises(ValueError):render_pdf_page(content,3)
    pdf.close()


def test_docx_bulk_import_real_doc():
    from docx import Document
    doc=Document()
    doc.add_paragraph('Jane Candidate completed Bachelor of Computer Science in 2019.')
    doc.add_paragraph('Backend Engineer 2020 to 2025 using Node.js and PostgreSQL with Docker deployment.')
    b=io.BytesIO();doc.save(b)
    result=ensure_valid_cv('jane.docx',b.getvalue())
    assert 'PostgreSQL' in result.text


def test_fictional_demo_dataset_ranks_suitable_candidate_first():
    from pathlib import Path
    folder=Path(__file__).resolve().parents[1]/'examples'
    jd=(folder/'Demo_Backend_Job.txt').read_text()
    cv=[(f.name,f.read_bytes()) for f in folder.glob('*.txt') if 'Job' not in f.name]
    result=evaluate_batch(jd,cv)
    assert len(result.candidates)==4
    assert result.ranked[0].name=='Jane_Backend.txt'
    assert any(x.status=='rejected' for x in result.candidates)
    assert all(x.score>=0 for x in result.candidates)
